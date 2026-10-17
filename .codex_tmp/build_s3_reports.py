from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1] / 'product-ongoing-real-data'
OUT = ROOT / 'deliverables'
OUT.mkdir(exist_ok=True)


def make_report(filename, title, task, opening, changes, outcome, limits):
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin, section.bottom_margin = Cm(1.8), Cm(1.8)
    section.left_margin, section.right_margin = Cm(2), Cm(2)
    normal = doc.styles['Normal']
    normal.font.name, normal.font.size = 'Arial', Pt(11)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.line_spacing = 1.12
    normal.paragraph_format.space_after = Pt(6)
    for name, size in [('Title', 18), ('Heading 1', 12), ('List Bullet', 11)]:
        style = doc.styles[name]
        style.font.name, style.font.size = 'Arial', Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_after = Pt(6)
        if name == 'Heading 1':
            style.font.bold = True
            style.paragraph_format.space_before = Pt(10)
    doc.core_properties.title = title
    doc.core_properties.subject = task
    doc.core_properties.author = ''
    doc.core_properties.keywords = 'TB-INSIGHT; ' + task
    doc.add_paragraph(title, 'Title')
    p = doc.add_paragraph()
    p.add_run(task).bold = True
    p.add_run(' | Penjelasan perubahan per 15 Oktober 2026')
    doc.add_paragraph(opening)
    doc.add_heading('Perubahan yang telah diterapkan', 1)
    for heading, text in changes:
        p = doc.add_paragraph(style='List Bullet')
        p.add_run(heading + '. ').bold = True
        p.add_run(text)
    doc.add_heading('Hasil dan pemeriksaan', 1)
    doc.add_paragraph(outcome)
    doc.add_heading('Catatan penggunaan', 1)
    doc.add_paragraph(limits)
    p = doc.add_paragraph('Acuan: rencana sprint TB-INSIGHT, API–UI Data Contract, protokol proyek, dan implementasi pada product-ongoing-real-data/frontend/.')
    for run in p.runs:
        run.font.size = Pt(9)
    # Keep the document language available to Word's proofing and accessibility tools.
    lang = OxmlElement('w:lang')
    lang.set(qn('w:val'), 'id-ID')
    normal.element.get_or_add_rPr().append(lang)
    for root in [doc.element, doc.styles.element]:
        for border in list(root.xpath('.//w:pBdr')):
            border.getparent().remove(border)
    target = OUT / filename
    doc.save(target)
    print(target)


make_report(
    'S3_ERI_01_Penjelasan_Perubahan.docx',
    'Penjelasan Perubahan Frontend Alpha TB INSIGHT',
    'S3-ERI-01',
    'Frontend alpha telah dikembangkan untuk mengeksplorasi tren pengobatan TB dan perbedaan antarwilayah. Dashboard menggunakan React dan Vite, terhubung ke backend FastAPI, serta memakai snapshot data resmi proyek.',
    [
        ('Navigasi halaman', 'Ringkasan, Tren, Peta, Provinsi, Metode, dan Sumber dipisahkan menjadi halaman yang dapat dibuka melalui menu.'),
        ('Filter bersama', 'Pilihan wilayah, rentang bulan, bulan fokus, dan jumlah atau rate disimpan dalam URL dan dibawa ketika pengguna berpindah halaman.'),
        ('Grafik tren', 'Grafik bulanan dilengkapi tooltip, penanda provisional, pembanding maksimal tiga provinsi, kalender, tabel, dan tampilan besar.'),
        ('Peta interaktif', 'Peta menampilkan indikator TB, konteks sosial, iklim, bencana, serta lampiran LISA dan persistence. Tersedia legenda, navigasi bulan, putar atau jeda, zoom, dan pembukaan profil provinsi.'),
        ('Profil dan unduhan', 'Profil wilayah memuat outcome, konteks sosial, iklim, bencana, serta perbandingan nasional. Data dan metadata dapat diunduh untuk ditelusuri kembali.'),
        ('Keterbacaan dan status data', 'Latar Azure Mist, font yang lebih besar, serta tampilan responsif diterapkan. Nol, nilai hilang, nilai tidak valid, dan data provisional dibedakan.'),
    ],
    'Dashboard memuat 912 observasi pada 38 provinsi selama 24 bulan, Januari 2024 sampai Desember 2025. Build frontend berhasil dan rangkaian 18 tes otomatis proyek lulus, termasuk agregasi, skala peta, status data, dan provenance.',
    'Angka menunjukkan pasien mulai pengobatan TBC sensitif obat, bukan estimasi insidensi. Hasil spasial Dame masih dipisahkan dari kohort aktif sampai rekonsiliasi input selesai. Visual tanpa data menampilkan kebutuhan data yang belum tersedia.'
)

make_report(
    'S3_ERI_02_Penjelasan_Perubahan.docx',
    'Penjelasan Perubahan Sumber dan Metode TB INSIGHT',
    'S3-ERI-02',
    'Halaman Sumber, Metode, dan catatan pada visual utama telah dilengkapi agar pengguna dapat menelusuri asal angka, memahami perhitungan, dan membaca keterbatasannya langsung di dalam dashboard.',
    [
        ('Informasi sumber', 'Setiap entri memuat penerbit, tautan asli, cakupan, satuan, peran, pemakaian pada visual, keterbatasan, dan jejak berkas. Input resmi, referensi pembanding, dan lampiran tim dibedakan.'),
        ('Pencatatan tanggal', 'Tanggal akses historis dipisahkan dari waktu yang tercatat pada olahan dan waktu pemeriksaan tautan. Tanggal tanpa bukti ditampilkan sebagai Belum tercatat.'),
        ('Penjelasan metode', 'Kamus indikator dan istilah ditambahkan, termasuk definisi outcome, rumus rate, pending outcome, provisional, data kosong, dan batas penggunaan hasil spasial.'),
        ('Catatan visual', 'Grafik, peta, dan ringkasan memuat sumber, satuan, periode, serta peringatan yang relevan. Pengguna dapat membuka definisi dan menelusuri sumber khusus visual tersebut.'),
        ('Penelusuran dan unduhan', 'Halaman Sumber memiliki pencarian, penyaringan jenis sumber, informasi pemakaian, checksum, serta unduhan daftar sumber dan metadata. Versi dokumentasi dibedakan dari versi data.'),
    ],
    'Metadata mencakup enam entri ledger dan satu lampiran analisis tim, dengan sembilan definisi indikator. Dari delapan tautan referensi yang diperiksa, lima dapat diakses dan tiga belum dapat diverifikasi. Seluruh 18 tes otomatis proyek dan build frontend berhasil.',
    'Tanggal akses dataset historis yang tidak tersimpan belum dapat dipulihkan. Pemeriksaan tautan tidak memperbarui angka penelitian. Persetujuan metode oleh reviewer, aturan kematangan, rekonsiliasi input Dame, dan manifest freeze tetap memerlukan tindak lanjut tim.'
)
