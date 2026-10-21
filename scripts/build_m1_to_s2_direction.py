from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "product-ongoing-real-data" / "deliverables" / "TB_INSIGHT_ClickUp_Tasks_64_SINGLE_SHEET.xlsx"
OUTPUT = ROOT / "outputs" / "2026-09-02-m1-to-s2" / "TB_INSIGHT_M1_to_S2_Direction.xlsx"

UPSTREAM = {
    "S2-VEN-01": ["S1-VEN-02", "S1-STE-01"],
    "S2-VEN-02": ["S1-VEN-02", "S1-STE-01", "S1-DAM-01"],
    "S2-STE-01": ["S1-VEN-02", "S1-STE-01", "S1-STE-02"],
    "S2-STE-02": ["S2-STE-01"],
    "S2-DAM-01": ["S1-VEN-02", "S1-DAM-01", "S1-DAM-02"],
    "S2-DAM-02": ["S2-DAM-01"],
    "S2-ERI-01": ["S1-ERI-01", "S1-ERI-02", "S1-VEN-02", "S2-VEN-01"],
    "S2-ERI-02": ["S2-ERI-01", "S1-ERI-02", "S1-VEN-02"],
    "S2-ADZ-01": ["S1-ADZ-01", "S1-ADZ-02", "S1-STE-01", "S1-DAM-02", "S1-VEN-02"],
    "S2-ADZ-02": ["S2-VEN-01", "S2-VEN-02", "S1-ADZ-02"],
}


def read_tasks():
    wb = load_workbook(SOURCE, read_only=True, data_only=True)
    ws = wb["ClickUp Import"]
    rows = list(ws.iter_rows(values_only=True))
    headers = {name: idx for idx, name in enumerate(rows[0])}
    return {
        row[headers["Task ID"]]: {name: row[idx] for name, idx in headers.items()}
        for row in rows[1:]
        if row[headers["Task ID"]]
    }


def joined_deliverables(ids, tasks):
    items = []
    for task_id in ids:
        task = tasks.get(task_id)
        if task:
            items.append(f"{task_id}: {task['Deliverable Path']}")
        else:
            items.append(f"{task_id}: output task S2 sebelumnya")
    return "\n".join(items)


def style_table(ws, header_row, last_row, widths):
    navy = "17365D"
    pale_blue = "DCE6F1"
    border = Border(bottom=Side(style="thin", color="B4C6E7"))
    for cell in ws[header_row]:
        cell.fill = PatternFill("solid", fgColor=navy)
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in ws.iter_rows(min_row=header_row + 1, max_row=last_row):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = border
        if row[0].row % 2 == 0:
            for cell in row:
                cell.fill = PatternFill("solid", fgColor="F7F9FC")
    for idx, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.auto_filter.ref = f"A{header_row}:{get_column_letter(len(widths))}{last_row}"
    ws.freeze_panes = f"A{header_row + 1}"
    ws.sheet_view.showGridLines = False


def build():
    tasks = read_tasks()
    wb = Workbook()
    ws = wb.active
    ws.title = "S2 Direction"

    ws.merge_cells("A1:H1")
    ws["A1"] = "TB-INSIGHT — M1 to S2 Direction"
    ws["A1"].fill = PatternFill("solid", fgColor="17365D")
    ws["A1"].font = Font(color="FFFFFF", bold=True, size=16)
    ws["A1"].alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 30

    ws.merge_cells("A2:H2")
    ws["A2"] = (
        "S1 telah selesai. Task tanpa dependency S2 dapat masuk READY setelah file S1 diverifikasi; task yang "
        "menunggu task S2 lain tetap BACKLOG. M1 tetap menjadi gate, bukan pengganti dependency langsung."
    )
    ws["A2"].fill = PatternFill("solid", fgColor="E2F0D9")
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[2].height = 42

    headers = [
        "S2 Task ID",
        "Task S2",
        "Owner",
        "Waiting On",
        "Required Input Files",
        "Arahan Pengerjaan",
        "Output S2",
        "Start Status",
    ]
    for col, value in enumerate(headers, 1):
        ws.cell(4, col, value)

    row_no = 5
    for task_id in UPSTREAM:
        task = tasks[task_id]
        deps = UPSTREAM[task_id]
        has_s2_dependency = any(dep.startswith("S2-") for dep in deps)
        start_status = "BACKLOG — wait for S2 upstream" if has_s2_dependency else "READY — verify S1 files"
        values = [
            task_id,
            task["Task Name"],
            task["Assignee"],
            ", ".join(deps),
            joined_deliverables(deps, tasks),
            task["Description"],
            task["Deliverable Path"],
            start_status,
        ]
        for col, value in enumerate(values, 1):
            ws.cell(row_no, col, value)
        if has_s2_dependency:
            ws.cell(row_no, 8).fill = PatternFill("solid", fgColor="FCE4D6")
            ws.cell(row_no, 8).font = Font(bold=True, color="9C5700")
        else:
            ws.cell(row_no, 8).fill = PatternFill("solid", fgColor="C6E0B4")
            ws.cell(row_no, 8).font = Font(bold=True, color="375623")
        row_no += 1

    style_table(ws, 4, row_no - 1, [14, 38, 13, 32, 48, 48, 38, 24])
    for row in range(5, row_no):
        ws.row_dimensions[row].height = 78

    inputs = wb.create_sheet("S1 Input Files")
    inputs.merge_cells("A1:F1")
    inputs["A1"] = "S1 Deliverables Required by S2"
    inputs["A1"].fill = PatternFill("solid", fgColor="17365D")
    inputs["A1"].font = Font(color="FFFFFF", bold=True, size=15)
    inputs["A1"].alignment = Alignment(vertical="center")
    inputs.row_dimensions[1].height = 28

    input_headers = ["S1 Task ID", "Task S1", "Owner", "Final Deliverable", "Used By", "Check Before S2"]
    for col, value in enumerate(input_headers, 1):
        inputs.cell(3, col, value)

    required_s1 = sorted({dep for deps in UPSTREAM.values() for dep in deps if dep.startswith("S1-")})
    row_no = 4
    for task_id in required_s1:
        task = tasks[task_id]
        used_by = [s2 for s2, deps in UPSTREAM.items() if task_id in deps]
        values = [
            task_id,
            task["Task Name"],
            task["Assignee"],
            task["Deliverable Path"],
            ", ".join(used_by),
            "COMPLETE — attach/link final reviewed file",
        ]
        for col, value in enumerate(values, 1):
            inputs.cell(row_no, col, value)
        inputs.cell(row_no, 6).fill = PatternFill("solid", fgColor="C6E0B4")
        inputs.cell(row_no, 6).font = Font(bold=True, color="375623")
        row_no += 1

    style_table(inputs, 3, row_no - 1, [14, 42, 14, 46, 32, 38])
    for row in range(4, row_no):
        inputs.row_dimensions[row].height = 48

    wb.properties.title = "TB-INSIGHT M1 to S2 Direction"
    wb.properties.subject = "Direct upstream tasks and required S1 deliverables for Sprint 2"
    wb.properties.creator = "TB-INSIGHT"
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build())
